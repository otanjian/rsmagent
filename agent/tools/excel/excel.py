# encoding:utf-8
"""Excel tool - inspect and edit an .xlsx workbook, natively.

Why this exists: the ``read`` tool renders a workbook as text, which answers
"what does it say" but not "which column is 文档索引" or "clear column J on
every data row". Those two questions used to send the model to ``bash`` to
hand-write a Python or PowerShell script, and on the deployment this was
prompted by there is no MS Office COM, no ACE.OLEDB driver and (before the
launcher fix) no ``python`` on PATH - so the script failed, the turn was spent,
and the user got a menu of workarounds instead of an answer.

Two actions, and that is the whole surface:

* ``inspect`` - sheets, real used range, the candidate header rows and the
  column names in them, plus an optional row dump. This is what answers "which
  column is X" without asking the user to open the file.
* ``update`` - set or clear one column over a row range, addressed by header
  **name** or by column letter. Header resolution is the part that used to
  require the model to guess.

Both are openpyxl calls, so they work wherever the app runs: no Excel install,
no COM, no OleDb. Formulas, number formats and styling survive an update because
the workbook is loaded with ``data_only=False`` (a value-only load would replace
every formula with its cached result and then save that). Embedded images and
charts do not survive: openpyxl does not round-trip drawings, so a workbook that
carries them gets a warning in the result rather than silent loss.
"""

import os
import re
import zipfile
from typing import Any, Dict, List, Optional, Tuple

from agent.tools.base_tool import BaseTool, ToolResult
from agent.tools.utils.file_state import note_read, note_write, staleness_warning
from agent.tools.utils.truncate import truncate_head, format_size, DEFAULT_MAX_BYTES
from common.log import logger
from common.utils import expand_path


#: Rows searched for a header row when ``header_row`` is not given. Real
#: templates put the table below a title block, so row 1 alone is not enough.
HEADER_SCAN_ROWS = 20

#: Cells printed per row / characters printed per cell in an inspect report.
MAX_DISPLAY_COLS = 40
CELL_DISPLAY_CHARS = 60

#: Rows dumped by one inspect call, and the ceiling on cells one update may
#: touch (a mistaken "all rows" on a bloated sheet should not run for minutes).
MAX_SAMPLE_ROWS = 200
MAX_UPDATE_CELLS = 200000

SHEET_FORMATS = ('.xlsx', '.xlsm', '.xltx', '.xltm')


def _normalize_header(value: Any) -> str:
    """Header text reduced to what a comparison can use.

    Template headers carry newlines, full-width spaces and parenthesised
    qualifications ("文档索引\\n(可留空)"), and the model quotes them from
    memory, so whitespace is stripped everywhere rather than trimmed at the ends.
    """
    text = str(value) if value is not None else ""
    return re.sub(r"\s+", "", text).casefold()


def _display(value: Any) -> str:
    """One cell as a single line, bounded in length."""
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\n", " ").strip()
    if len(text) > CELL_DISPLAY_CHARS:
        text = text[:CELL_DISPLAY_CHARS] + "..."
    return text


def _used_bounds(ws) -> Tuple[int, int, int, int]:
    """(min_row, max_row, min_col, max_col) of the cells that exist.

    Read from the cell table instead of ``ws.max_row`` on purpose: a workbook
    saved by another tool can carry a dimension record of a million rows, and
    iterating or trusting that number turns a 53-row sheet into a minute of
    work. ``_cells`` is exactly the set of cells the file actually contains.
    """
    cells = getattr(ws, "_cells", None)
    if cells:
        rows = [row for row, _col in cells]
        cols = [col for _row, col in cells]
        return min(rows), max(rows), min(cols), max(cols)
    # read-only worksheets have no cell table; their dimension is all we have.
    return 1, max(ws.max_row or 1, 1), 1, max(ws.max_column or 1, 1)


def _cell_value(ws, row: int, col: int):
    """Read one cell without creating it (``ws.cell`` materialises empty cells)."""
    cells = getattr(ws, "_cells", None)
    if cells is not None:
        cell = cells.get((row, col))
        return None if cell is None else cell.value
    return ws.cell(row=row, column=col).value


def _column_index(text: str) -> Optional[int]:
    """``"J"`` / ``"j"`` -> 10, or None when the text is not a column letter."""
    candidate = str(text or "").strip().upper()
    if not re.fullmatch(r"[A-Z]{1,3}", candidate):
        return None
    try:
        from openpyxl.utils import column_index_from_string

        return column_index_from_string(candidate)
    except Exception:
        return None


def _parse_rows(spec: Any, first: int, last: int) -> Tuple[Optional[List[int]], Optional[str]]:
    """Parse a row selection into an explicit list, or explain why not.

    Accepts ``"2-53"``, ``"2-"``, ``"-53"``, ``"5"``, ``"2-10,15,20-25"`` and
    ``"all"``. Ranges are interpreted against the sheet (``first``/``last`` are
    its used bounds), never against the file's declared dimension.
    """
    if first > last:
        return [], None
    text = str(spec or "").strip().lower()
    if text in ("", "all", "*"):
        rows = list(range(first, last + 1))
    else:
        rows = []
        for part in text.split(","):
            part = part.strip()
            if not part:
                continue
            match = re.fullmatch(r"(\d*)\s*-\s*(\d*)", part)
            if match and (match.group(1) or match.group(2)):
                start = int(match.group(1)) if match.group(1) else first
                end = int(match.group(2)) if match.group(2) else last
            elif re.fullmatch(r"\d+", part):
                start = end = int(part)
            else:
                return None, (
                    f"'{part}' is not a row range. Use \"2-53\", \"2-\", \"5\", "
                    f"\"2-10,15\" or \"all\"."
                )
            if start < 1 or end < start:
                return None, f"'{part}' is not a valid row range (rows start at 1)."
            rows.extend(range(start, end + 1))
    if not rows:
        return None, "the row selection is empty."
    rows = sorted(set(rows))
    if len(rows) > MAX_UPDATE_CELLS:
        return None, (
            f"the selection covers {len(rows)} rows, more than the "
            f"{MAX_UPDATE_CELLS} cell limit of one call. Narrow it down."
        )
    return rows, None


def _header_rows(ws, max_col: int, limit: int) -> List[Tuple[int, List[Tuple[int, Any]]]]:
    """The first ``limit`` rows that carry at least two named cells."""
    out: List[Tuple[int, List[Tuple[int, Any]]]] = []
    _, max_row, _, _ = _used_bounds(ws)
    for row in range(1, min(max_row, HEADER_SCAN_ROWS) + 1):
        cells = [
            (col, _cell_value(ws, row, col))
            for col in range(1, min(max_col, MAX_DISPLAY_COLS) + 1)
        ]
        filled = [(col, value) for col, value in cells if value not in (None, "")]
        if len(filled) >= 2 or (filled and not out):
            out.append((row, filled))
        if len(out) >= limit:
            break
    return out


def _resolve_column(
    ws, name: str, max_col: int, header_row: Optional[int] = None
) -> Tuple[Optional[int], Optional[int], Optional[str], Optional[str]]:
    """Find the column whose header matches ``name``.

    Returns ``(column_index, header_row, header_text, error)``. Matching is
    exact after normalization, then prefix/suffix, then substring - so
    "文档索引" finds "文档索引(可留空)" and stays unambiguous when both exist.
    The error names the header rows it saw, which is what turns "I don't know
    the column name" into a self-correcting second call.
    """
    target = _normalize_header(name)
    if not target:
        return None, None, None, "column is required (a header name or a column letter)."

    rows = [header_row] if header_row else list(range(1, HEADER_SCAN_ROWS + 1))
    best: Optional[Tuple[int, int, int, Any]] = None  # score, row, col, text
    for row in rows:
        for col in range(1, min(max_col, MAX_DISPLAY_COLS) + 1):
            value = _cell_value(ws, row, col)
            norm = _normalize_header(value)
            if not norm:
                continue
            if norm == target:
                score = 3
            elif norm.startswith(target) or norm.endswith(target):
                score = 2
            elif target in norm:
                score = 1
            else:
                continue
            if best is None or score > best[0] or (score == best[0] and row < best[1]):
                best = (score, row, col, value)

    if best is not None:
        score, row, col, value = best
        if score < 3:
            # A partial match is used, but the model is told which one it got:
            # on a template with both "文档索引" and "文档索引(补充)" a silent
            # prefix match is how the wrong column gets cleared.
            logger.debug(
                "[Excel] column %r matched %r on row %s (partial)", name, value, row
            )
        return col, row, str(value), None

    return None, None, None, "no column header matched %r." % name


def _header_hint(ws, max_col: int) -> str:
    """The header rows and their cells, printed for a failed column lookup."""
    lines = []
    for row, filled in _header_rows(ws, max_col, limit=3):
        if not filled:
            continue
        rendered = ", ".join(f"{_letter(col)}={_display(value)}" for col, value in filled)
        lines.append(f"row {row}: {rendered}")
    if not lines:
        return "The sheet has no header row in its first %d rows." % HEADER_SCAN_ROWS
    return "Header rows found:\n" + "\n".join(lines)


def _letter(col: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(col)


def _has_drawings(path: str) -> bool:
    """Whether the workbook embeds images or charts (lost by an openpyxl save)."""
    try:
        with zipfile.ZipFile(path) as archive:
            return any(
                name.startswith(("xl/media/", "xl/drawings/", "xl/charts/"))
                for name in archive.namelist()
            )
    except Exception:
        return False


def _load(absolute_path: str, data_only: bool):
    """Load a workbook, or raise ImportError/ValueError with a usable message."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportError(
            "Error: openpyxl library not installed. Install with: pip install openpyxl"
        )
    keep_vba = os.path.splitext(absolute_path)[1].lower() in ('.xlsm', '.xltm')
    return load_workbook(absolute_path, data_only=data_only, keep_vba=keep_vba)


class Excel(BaseTool):
    """Inspect and edit .xlsx workbooks without writing a script."""

    name: str = "excel"

    description: str = (
        "Inspect or edit an Excel workbook (.xlsx/.xlsm) directly. Use this instead of writing a "
        "Python/PowerShell script, and instead of Excel COM / OleDb automation - the server runs "
        "headless with no MS Office installed. "
        "action='inspect' reports the sheets, the real used range, the header candidate rows and "
        "the column names in them (that is how you learn the exact column name to address); it never "
        "modifies the file. "
        "action='update' sets or clears ONE column over a row range, addressing the column by its "
        "header name (e.g. \"文档索引\") or by letter (e.g. \"J\"). Omit `value` to clear the cells; "
        "pass a value to write it. Only the named column and rows are touched: formulas, styling and "
        "every other cell are preserved. Quote header names from the inspect output rather than from "
        "memory, and inspect first when you are not sure which column or sheet is meant."
    )

    params: dict = {
        "type": "object",
        "properties": {
            "action": {
                "type": "string",
                "enum": ["inspect", "update"],
                "description": "'inspect' to read the structure (sheets, header rows, column names, optional row samples); 'update' to set or clear one column over a row range."
            },
            "path": {
                "type": "string",
                "description": "Workbook path (.xlsx/.xlsm). Relative paths resolve against the workspace directory; use an absolute path (or ~/...) for files elsewhere."
            },
            "sheet": {
                "type": "string",
                "description": "Sheet name. Optional when the workbook has exactly one sheet; for a multi-sheet workbook 'inspect' lists all of them and 'update' requires the name."
            },
            "column": {
                "type": "string",
                "description": "update only: the column to write, as a header name (\"文档索引\") or a column letter (\"J\")."
            },
            "rows": {
                "type": "string",
                "description": "update only: rows to change, e.g. \"2-53\", \"2-\", \"5\", \"2-10,15\", or \"all\" for every data row. Defaults to the data rows under the resolved header row. 'inspect' also accepts it, to dump that range."
            },
            "value": {
                "type": "string",
                "description": "update only: the text to write into every selected cell. Omit it (or pass an empty string) to clear the cells instead. A numeric-looking value is written as text so leading zeros survive."
            },
            "header_row": {
                "type": "integer",
                "description": "Row number of the header row, when it is known (inspect reports it). Auto-detected within the first 20 rows when omitted."
            },
            "output": {
                "type": "string",
                "description": "update only: save to this path instead of overwriting the source file. Use it when the original template must stay untouched."
            },
            "header_rows": {
                "type": "integer",
                "description": "inspect only: how many leading rows to print (default 6)."
            },
            "sample_rows": {
                "type": "string",
                "description": "inspect only: dump this row range, e.g. \"1-10\"."
            }
        },
        "required": ["action", "path"]
    }

    def __init__(self, config: dict = None):
        self.config = config or {}
        self.cwd = self.config.get("cwd", os.getcwd())

    # ------------------------------------------------------------------
    # entry point
    # ------------------------------------------------------------------

    def execute(self, args: Dict[str, Any]) -> ToolResult:
        action = str(args.get("action") or "").strip().lower()
        path = str(args.get("path") or "").strip()
        if not path:
            return ToolResult.fail("Error: path is required.")
        if action not in ("inspect", "update"):
            return ToolResult.fail(
                "Error: action must be 'inspect' or 'update' (got %r)." % action
            )

        absolute_path = self._resolve_path(path)
        if not os.path.exists(absolute_path):
            hint = ""
            if not os.path.isabs(path) and not path.startswith("~"):
                hint = (
                    f"\nResolved to: {absolute_path}"
                    f"\nHint: relative paths are based on the workspace ({self.cwd}); "
                    f"use an absolute path for files outside it."
                )
            return ToolResult.fail(f"Error: path not found: {path}{hint}")
        if os.path.isdir(absolute_path):
            return ToolResult.fail(f"Error: {path} is a directory, not a workbook.")

        suffix = os.path.splitext(absolute_path)[1].lower()
        if suffix == '.xls':
            return ToolResult.fail(
                "Error: the legacy .xls format cannot be read or written (openpyxl "
                "handles .xlsx/.xlsm only). Ask the user to save it as .xlsx."
            )
        if suffix not in SHEET_FORMATS:
            return ToolResult.fail(
                f"Error: {suffix or 'this file'} is not an Excel workbook. Expected "
                f"{', '.join(SHEET_FORMATS)}. Use the read tool for other formats."
            )

        display = self._display_path(path, absolute_path)
        try:
            if action == "inspect":
                return self._inspect(absolute_path, display, args)
            return self._update(absolute_path, display, args)
        except ImportError as e:
            return ToolResult.fail(str(e))
        except PermissionError:
            return ToolResult.fail(
                f"Error: permission denied writing {path}. On Windows this usually "
                f"means the file is open in Excel - ask the user to close it and retry."
            )
        except Exception as e:
            logger.error("[Excel] %s failed on %s: %s", action, absolute_path, e)
            return ToolResult.fail(f"Error: could not {action} {path}: {e}")

    # ------------------------------------------------------------------
    # inspect
    # ------------------------------------------------------------------

    def _inspect(self, absolute_path: str, display: str, args: Dict[str, Any]) -> ToolResult:
        workbook = _load(absolute_path, data_only=True)
        try:
            names = list(workbook.sheetnames)
            wanted = str(args.get("sheet") or "").strip()
            if wanted and wanted not in names:
                close = [name for name in names if wanted.casefold() in name.casefold()]
                extra = f" Did you mean {close!r}?" if close else ""
                return ToolResult.fail(
                    f"Error: no sheet named {wanted!r} in {display}. Sheets: {names}.{extra}"
                )
            selected = [wanted] if wanted else names

            header_rows = args.get("header_rows")
            try:
                header_rows = int(header_rows) if header_rows is not None else 6
            except (TypeError, ValueError):
                header_rows = 6

            sample = str(args.get("sample_rows") or "").strip()

            lines = [
                f"file: {display}",
                f"size: {format_size(os.path.getsize(absolute_path))}",
                f"sheets ({len(names)}): {', '.join(names)}",
            ]
            for name in selected:
                ws = workbook[name]
                lines.extend(self._sheet_report(ws, name, header_rows, sample))

            content = "\n".join(lines)
            truncation = truncate_head(content, max_lines=2000, max_bytes=DEFAULT_MAX_BYTES)
            output = truncation.content
            if truncation.truncated:
                output += (
                    f"\n\n[report truncated at {format_size(DEFAULT_MAX_BYTES)}; "
                    f"inspect one sheet at a time with sheet=<name>]"
                )
            note_read(absolute_path)
            return ToolResult.success({
                "output": output,
                "path": display,
                "sheets": names,
                "details": {"truncation": truncation.to_dict()} if truncation.truncated else None,
            })
        finally:
            try:
                workbook.close()
            except Exception:
                pass

    def _sheet_report(self, ws, name: str, header_rows: int, sample: str) -> List[str]:
        min_row, max_row, min_col, max_col = _used_bounds(ws)
        lines = [
            "",
            f"=== sheet {name!r} | rows {min_row}-{max_row} | cols "
            f"{_letter(min_col)}-{_letter(max_col)} | {max_row - min_row + 1} rows x "
            f"{max_col - min_col + 1} cols ===",
        ]

        leading = _header_rows(ws, max_col, limit=header_rows)
        for row, filled in leading:
            if not filled:
                continue
            rendered = " | ".join(f"{_letter(col)}={_display(value)}" for col, value in filled)
            lines.append(f"r{row}: {rendered}")

        # The column names of the row that looks most like a header, one per
        # line: this is the list the model should quote back in `column`. A row
        # is only claimed as a header when it names at least three columns - a
        # key-value block ("公司 | 恒信精密…") has rows of two named cells and is
        # not a table, and calling one of those the header sends the model to
        # the wrong column with confidence.
        candidates = [
            (row, filled) for row, filled in leading
            if sum(1 for _col, value in filled if _normalize_header(value)) >= 3
        ]
        if candidates:
            best_row, filled = max(
                candidates, key=lambda item: sum(1 for c, v in item[1] if _normalize_header(v))
            )
            lines.append(
                f"columns named in row {best_row} (address a column by its name or letter):"
            )
            for col, value in filled:
                if _normalize_header(value):
                    lines.append(f"  {_letter(col)} = {_display(value)}")
        else:
            lines.append(
                "[no table header row found in the leading rows - this may be a "
                "key-value layout; pass header_row=<row> if you know which row is the header]"
            )

        if sample:
            rows, error = _parse_rows(sample, min_row, max_row)
            if error:
                lines.append(f"[sample_rows ignored: {error}]")
            elif not rows:
                lines.append("[sample_rows ignored: the sheet has no rows to dump]")
            else:
                show = rows[:MAX_SAMPLE_ROWS]
                lines.append(f"rows {show[0]}-{show[-1]} ({len(show)} shown):")
                for row in show:
                    values = [
                        _display(_cell_value(ws, row, col))
                        for col in range(min_col, min(max_col, min_col + MAX_DISPLAY_COLS - 1) + 1)
                    ]
                    lines.append(f"r{row}: " + " | ".join(values).rstrip(" |"))
                if len(rows) > len(show):
                    lines.append(f"[{len(rows) - len(show)} more rows not shown]")
        else:
            lines.append(
                f"[dump rows with action='inspect', sheet={name!r}, sample_rows=\"1-10\"]"
            )
        return lines

    # ------------------------------------------------------------------
    # update
    # ------------------------------------------------------------------

    def _update(self, absolute_path: str, display: str, args: Dict[str, Any]) -> ToolResult:
        workbook = _load(absolute_path, data_only=False)
        try:
            names = list(workbook.sheetnames)
            wanted = str(args.get("sheet") or "").strip()
            if not wanted:
                if len(names) != 1:
                    return ToolResult.fail(
                        f"Error: this workbook has {len(names)} sheets ({names}); pass "
                        f"sheet=<name> so the update does not hit the wrong one."
                    )
                ws = workbook[names[0]]
            else:
                if wanted not in names:
                    return ToolResult.fail(
                        f"Error: no sheet named {wanted!r} in {display}. Sheets: {names}."
                    )
                ws = workbook[wanted]
            sheet_name = ws.title

            min_row, max_row, _min_col, max_col = _used_bounds(ws)
            if max_row < 1:
                return ToolResult.fail(
                    f"Error: sheet {sheet_name!r} is empty; there is nothing to update."
                )

            header_row = args.get("header_row")
            if header_row is not None:
                try:
                    header_row = int(header_row)
                except (TypeError, ValueError):
                    return ToolResult.fail("Error: header_row must be a row number.")
                if not 1 <= header_row <= max_row:
                    return ToolResult.fail(
                        f"Error: header_row {header_row} is outside the sheet "
                        f"(rows {min_row}-{max_row})."
                    )

            column_arg = str(args.get("column") or "").strip()
            if not column_arg:
                return ToolResult.fail(
                    "Error: update needs `column` (a header name or letter). Run "
                    f"action='inspect' on {display} first to see the column names."
                )

            col = _column_index(column_arg)
            if col is not None:
                resolved_header_row = header_row
                if resolved_header_row is None:
                    resolved_header_row = self._auto_header_row(ws, col)
                header_text = _display(_cell_value(ws, resolved_header_row or 1, col)) or None
            else:
                col, resolved_header_row, header_text, error = _resolve_column(
                    ws, column_arg, max_col, header_row
                )
                if error:
                    return ToolResult.fail(
                        f"Error: {error}\n{_header_hint(ws, max_col)}\n"
                        f"Pass `column` with one of those names, or a column letter."
                    )

            first_data_row = (resolved_header_row or min_row) + 1
            rows_arg = str(args.get("rows") or "").strip()
            if rows_arg.lower() in ("", "all", "*"):
                # "all" covers the data rows, not the header: the header is what
                # the column was resolved by, so writing into it is never meant.
                rows, error = _parse_rows(None, first_data_row, max_row)
            else:
                rows, error = _parse_rows(rows_arg, min_row, max_row)
            if error:
                return ToolResult.fail(f"Error: {error}")
            if not rows:
                return ToolResult.fail(
                    f"Error: the selection is empty - sheet {sheet_name!r} has no rows "
                    f"below row {resolved_header_row if resolved_header_row else max_row} "
                    f"to update. Pass an explicit `rows` range (it has {max_row} rows)."
                )

            if "value" in args and args.get("value") is not None:
                raw = args.get("value")
                value: Any = raw if isinstance(raw, (int, float, bool)) else str(raw)
                clearing = False
            else:
                value = None
                clearing = True

            changed = 0
            already = 0
            for row in rows:
                cell = ws.cell(row=row, column=col)
                if cell.value == value:
                    already += 1
                    continue
                cell.value = value
                changed += 1

            output_path = absolute_path
            output_arg = str(args.get("output") or "").strip()
            if output_arg:
                output_path = self._resolve_path(output_arg)
                parent = os.path.dirname(output_path)
                if parent and not os.path.isdir(parent):
                    os.makedirs(parent, exist_ok=True)

            warning = staleness_warning(output_path) if output_path == absolute_path else None
            if _has_drawings(absolute_path):
                logger.info("[Excel] %s embeds drawings/charts", absolute_path)

            workbook.save(output_path)
            note_write(output_path)

            where = (
                "in place" if output_path == absolute_path
                else f"to {self._display_path(output_arg, output_path)}"
            )
            header_label = f" ({header_text!r})" if header_text else ""
            action_label = "cleared" if clearing else f"set to {value!r}"
            lines = [
                f"updated {display} {where}",
                f"sheet {sheet_name!r}: column {_letter(col)}{header_label} - "
                f"{changed} cell(s) {action_label}"
                + (f", {already} already equal" if already else ""),
                f"rows: {rows[0]}-{rows[-1]} ({len(rows)} selected)"
                + (f" | header row {resolved_header_row}" if resolved_header_row else ""),
                f"file: {format_size(os.path.getsize(output_path))}",
            ]
            if _has_drawings(absolute_path):
                lines.append(
                    "note: this workbook embeds images/charts, which openpyxl cannot "
                    "round-trip - check the result if they matter."
                )
            if warning:
                lines.append(f"note: {warning}")

            return ToolResult.success({
                "output": "\n".join(lines),
                "path": self._display_path(output_arg, output_path) if output_arg else display,
                "sheet": sheet_name,
                "column": _letter(col),
                "rows": [rows[0], rows[-1]],
                "changed": changed,
            })
        finally:
            try:
                workbook.close()
            except Exception:
                pass

    def _auto_header_row(self, ws, col: int) -> Optional[int]:
        """The header row of one column, best effort, for the default row range.

        Only used when the caller addressed the column by letter: the header row
        is what "the data rows under the header" defaults to. The heuristic is
        the first of two adjacent non-empty cells in that column (a header sits
        above at least one value), falling back to the first non-empty row. The
        row it picked is printed in the result, so a wrong guess is visible and
        correctable with an explicit ``rows``.
        """
        first_filled: Optional[int] = None
        for row in range(1, HEADER_SCAN_ROWS + 1):
            value = _cell_value(ws, row, col)
            if value in (None, ""):
                continue
            if first_filled is None:
                first_filled = row
            below = _cell_value(ws, row + 1, col)
            if below not in (None, ""):
                return row
        return first_filled

    # ------------------------------------------------------------------
    # paths
    # ------------------------------------------------------------------

    def _resolve_path(self, path: str) -> str:
        path = expand_path(path)
        if os.path.isabs(path):
            return path
        return os.path.abspath(os.path.join(self.cwd, path))

    def _display_path(self, original: str, absolute: str) -> str:
        """Relative paths are echoed relative; absolute ones stay absolute."""
        if not os.path.isabs(original) and not original.startswith("~"):
            return original
        return absolute
