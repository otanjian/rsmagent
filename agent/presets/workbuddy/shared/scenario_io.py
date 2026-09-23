#!/usr/bin/env python3
"""Local file IO for scenario agents; business judgements stay in each skill.

prepare copies demo inputs, inspect extracts evidence with locations, render
writes supplied analysis, and check verifies the resulting files. None of these
commands performs business analysis or turns example conclusions into facts.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import shutil
from datetime import date, datetime
from pathlib import Path


def dump(value, path):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str)+"\n", encoding="utf-8")


def prepare(skill_dir, destination):
    skill_dir, destination = Path(skill_dir).resolve(), Path(destination).resolve()
    scenario = json.loads((skill_dir/"scenario.json").read_text(encoding="utf-8"))
    source = skill_dir/"assets/demo"
    if not source.is_dir() or not any(source.iterdir()):
        raise ValueError("演示资料不存在")
    if destination == skill_dir or skill_dir in destination.parents:
        raise ValueError("运行目录不能覆盖技能资源")
    destination.mkdir(parents=True, exist_ok=False)
    try:
        shutil.copytree(source, destination/"输入参考文件")
        (destination/"outputs").mkdir()
        (destination/"任务.md").write_text("# 演示任务 · "+scenario["name"]+"\n\n"+scenario["default_task"]+"\n", encoding="utf-8")
        result = {"scene": scenario["slug"], "mode": "demo", "state": "inputs_prepared",
                  "notice": "仅已准备输入。须读取原始资料、执行场景规则并生成报告，不能把本清单当作业务结果。",
                  "task": str(destination/"任务.md"), "outputs": str(destination/"outputs"),
                  "files": [{"path": str(p), "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}
                            for p in sorted((destination/"输入参考文件").rglob("*")) if p.is_file()]}
        dump(result, destination/"run.json")
        return result
    except Exception:
        shutil.rmtree(destination)
        raise


def inspect_file(path):
    path = Path(path)
    result = {"file": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=False, data_only=False)
        cached = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            result["sheets"] = []
            for sheet in book:
                values = cached[sheet.title]
                rows = []
                for row in sheet:
                    cells = []
                    for cell in row:
                        if cell.value is None:
                            continue
                        item = {"cell": cell.coordinate, "value": cell.value}
                        if cell.data_type == "f":
                            item = {"cell": cell.coordinate, "formula": cell.value, "cached_value": values[cell.coordinate].value}
                        if cell.comment:
                            item["comment"] = cell.comment.text
                        cells.append(item)
                    if cells:
                        rows.append(cells)
                result["sheets"].append({"name": sheet.title, "merged_ranges": [str(r) for r in sheet.merged_cells.ranges], "rows": rows})
        finally:
            book.close()
            cached.close()
    elif suffix == ".docx":
        from docx import Document
        doc = Document(path)
        result["paragraphs"] = [{"paragraph": i+1, "text": p.text} for i, p in enumerate(doc.paragraphs) if p.text.strip()]
        result["tables"] = [{"table": i+1, "rows": [[c.text for c in row.cells] for row in t.rows]} for i, t in enumerate(doc.tables)]
        result["embedded_images"] = len(doc.inline_shapes)
    elif suffix == ".pdf":
        import pdfplumber
        with pdfplumber.open(path) as doc:
            result["pages"] = [{"page": i+1, "text": page.extract_text() or ""} for i, page in enumerate(doc.pages)]
        result["needs_visual_review"] = any(not p["text"].strip() for p in result["pages"])
    elif suffix in (".txt", ".md", ".csv", ".json"):
        result["lines"] = [{"line": i+1, "text": line} for i, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines())]
    else:
        raise ValueError("不支持自动提取的格式："+suffix+"；请使用专用读取工具")
    return result


def safe_name(name):
    if not isinstance(name, str) or not name.strip() or name in (".", "..") or any(c in name for c in ("/", "\\", "\0")):
        raise ValueError("文件名必须是单个文件名")
    return name


def render_workbook(spec, path):
    import xlsxwriter
    sheets = spec.get("sheets", [])
    if not sheets or any(not s.get("columns") or not s.get("rows") for s in sheets):
        raise ValueError("工作簿必须包含有表头与实际数据行的工作表")
    with xlsxwriter.Workbook(str(path), {"strings_to_formulas": False, "strings_to_urls": False}) as book:
        title_fmt = book.add_format({"bold": True, "font_size": 15, "font_color": "#1A2332"})
        head_fmt = book.add_format({"bold": True, "bg_color": "#2D6CDF", "font_color": "white", "text_wrap": True})
        body_fmt = book.add_format({"text_wrap": True, "valign": "top"})
        date_fmt = book.add_format({"num_format": "yyyy-mm-dd"})
        for table in sheets:
            ws = book.add_worksheet(table["name"])
            cols, rows = table["columns"], table["rows"]
            ws.write(0, 0, spec["title"], title_fmt)
            ws.write(1, 0, spec["basis"], body_fmt)
            ws.write_row(2, 0, cols, head_fmt)
            for ri, row in enumerate(rows, start=3):
                if len(row) != len(cols):
                    raise ValueError(f"{table['name']} 第 {ri+1} 行与列数不一致")
                for ci, value in enumerate(row):
                    if isinstance(value, dict):
                        if set(value) != {"formula", "value"} or not str(value["formula"]).startswith("=") or value["value"] is None:
                            raise ValueError("公式须提供 formula 与已独立核算的 value 缓存")
                        if not isinstance(value["value"], (int, float)) or not math.isfinite(value["value"]):
                            raise ValueError("公式缓存须是有限数值")
                        ws.write_formula(ri, ci, value["formula"], body_fmt, value["value"])
                    elif isinstance(value, (date, datetime)):
                        ws.write_datetime(ri, ci, value, date_fmt)
                    else:
                        ws.write(ri, ci, value, body_fmt)
            ws.freeze_panes(3, 0)
            ws.autofilter(2, 0, len(rows)+2, len(cols)-1)
            ws.hide_gridlines(2)
            ws.set_column(0, len(cols)-1, 24, body_fmt)
            ws.set_landscape()
            ws.fit_to_pages(1, 0)
            ws.repeat_rows(0, 2)


def render_document(spec, path):
    if not spec.get("paragraphs"):
        raise ValueError("文档须有实际正文")
    if path.suffix.lower() == ".docx":
        from docx import Document
        doc = Document()
        doc.add_heading(spec["title"], 0)
        doc.add_paragraph(spec["basis"])
        for text in spec["paragraphs"]:
            doc.add_paragraph(text)
        doc.save(path)
    else:
        from xml.sax.saxutils import escape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        style = ParagraphStyle("zh", fontName="STSong-Light", fontSize=11, leading=18, wordWrap="CJK")
        content = [Paragraph(escape(spec["title"]), style), Spacer(1, 12), Paragraph(escape(spec["basis"]), style), Spacer(1, 12)]
        for text in spec["paragraphs"]:
            content.extend([Paragraph(escape(text).replace("\n", "<br/>"), style), Spacer(1, 8)])
        SimpleDocTemplate(str(path)).build(content)


def render(spec, destination):
    """Render explicit model/code-produced findings, never auto-fill facts."""
    if not spec.get("title") or not spec.get("basis"):
        raise ValueError("必须注明报告标题与数据依据（演示/真实、期间、来源）")
    files = spec.get("files", [])
    if not files:
        raise ValueError("没有交付物")
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    paths = [destination/safe_name(f["filename"]) for f in files]
    if len(set(paths)) != len(paths) or any(p.exists() for p in paths):
        raise ValueError("输出文件重复或已存在；请选择新的运行目录")
    created = []
    try:
        for file, path in zip(files, paths):
            merged = {**spec, **file}
            created.append(path)
            suffix = path.suffix.lower()
            if suffix == ".xlsx":
                render_workbook(merged, path)
            elif suffix in (".docx", ".pdf"):
                render_document(merged, path)
            elif suffix in (".md", ".txt") and file.get("text", "").strip():
                path.write_text(spec["title"]+"\n"+spec["basis"]+"\n\n"+file["text"], encoding="utf-8")
            else:
                raise ValueError("不支持的格式或空正文："+path.name)
        return [check(p) for p in paths]
    except Exception:
        for p in created:
            p.unlink(missing_ok=True)
        raise


def check(path):
    path = Path(path)
    content = inspect_file(path)
    warnings = []
    if "sheets" in content:
        if not content["sheets"] or any(not s["rows"] for s in content["sheets"]):
            raise ValueError("存在空工作表："+str(path))
        errors = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#SPILL!"}
        for sheet in content["sheets"]:
            for row in sheet["rows"]:
                for cell in row:
                    value = cell.get("cached_value", cell.get("value"))
                    if isinstance(value, str) and value in errors:
                        raise ValueError(f"{path.name}/{sheet['name']}/{cell['cell']}: {value}")
                    if "formula" in cell and value is None:
                        warnings.append(f"{sheet['name']}!{cell['cell']} 公式没有缓存，须另行重算验收")
    elif "pages" in content and (not content["pages"] or content["needs_visual_review"]):
        warnings.append("PDF 有无文本页，须视觉检查")
    elif "paragraphs" in content and not content["paragraphs"] and not content["tables"]:
        raise ValueError("空 Word 文档")
    elif "lines" in content and not any(r["text"].strip() for r in content["lines"]):
        raise ValueError("空文本文件")
    return {"file": str(path), "sha256": content["sha256"], "readable": True,
            "warnings": warnings, "business_validated": False}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    prep = sub.add_parser("prepare")
    prep.add_argument("--output", required=True, type=Path)
    prep.add_argument("--skill-dir", type=Path, default=Path(__file__).resolve().parents[1])
    read = sub.add_parser("inspect")
    read.add_argument("file", type=Path)
    read.add_argument("--output", required=True, type=Path)
    write = sub.add_parser("render")
    write.add_argument("spec", type=Path)
    write.add_argument("--output", required=True, type=Path)
    verify = sub.add_parser("check")
    verify.add_argument("files", nargs="+", type=Path)
    args = parser.parse_args()
    if args.command == "prepare":
        result = prepare(args.skill_dir, args.output)
    elif args.command == "inspect":
        result = inspect_file(args.file)
        dump(result, args.output)
        result = {"extracted": str(args.output), "source": str(args.file)}
    elif args.command == "render":
        result = render(json.loads(args.spec.read_text(encoding="utf-8")), args.output)
    else:
        result = [check(p) for p in args.files]
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))


if __name__ == "__main__":
    main()
