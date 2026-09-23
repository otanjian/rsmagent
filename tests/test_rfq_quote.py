import copy
import importlib.util
import json
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
SPEC = importlib.util.spec_from_file_location("rfq_quote", ROOT/"skills/rfq-quote/scripts/quote.py")
rfq = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(rfq)


@pytest.fixture
def data():
    return json.loads((ROOT/"skills/rfq-quote/assets/demo/quote-input.json").read_text())


def test_demo_reconciles_each_tier_and_retains_unverified_evidence(data):
    result = rfq.calculate(data)
    assert result["cost"] == Decimal("980.25")
    assert result["manufacturing"] == sum(x["amount"] for x in result["lines"])
    for t in result["tiers"]:
        assert t["manufacturing"] == result["manufacturing"] + sum(x["amount"] for x in t["adjustments"])
        assert t["management"] == rfq.money(t["manufacturing"] * Decimal("0.08"))
        assert t["price"] == rfq.money(t["cost"] * (1+t["markup"]))
        assert t["profit"] == rfq.money((t["price"]-t["cost"]) * t["quantity"])
    assert result["history_cost"] is None
    assert "禁止发出" in result["status"]
    assert [r["level"] for r in result["risks"]] == ["A", "A", "B", "B", "C", "A"]
    assert result["risks"][-1]["exposure"] is None
    assert result["tiers"][1]["markup"] == result["tiers"][0]["markup"]
    assert result["tiers"][2]["profit"] > result["tiers"][0]["profit"]


@pytest.mark.parametrize("value", [0, -1, 1.01, None, "NaN", "Infinity", True])
def test_invalid_utilization_cannot_silently_make_a_quote(data, value):
    data["material"]["utilization"] = value
    with pytest.raises(ValueError): rfq.calculate(data)


def test_risk_thresholds_and_qualitative_override(data):
    cost = rfq.calculate(data)["cost"]
    base = data["risks"][0]
    data["risks"] = [dict(base, kind=kind, exposure=str(amount)) for kind, amount in [
        ("other", "0"), ("other", cost*Decimal(".02")), ("other", cost*Decimal(".06")),
        ("welding", "1"), ("surface", "1"), ("other", "-100")]]
    assert [x["level"] for x in rfq.calculate(data)["risks"]] == ["C", "B", "A", "A", "B", "A"]
    data["risks"][0]["resolved"] = True
    with pytest.raises(ValueError, match="确认"): rfq.calculate(data)


def test_history_uses_adjusted_actual_cost_and_detects_large_deviation(data):
    data["history"] = {"actual_cost": 1000, "source": "ERP实际成本", "comparability": "同结构同工艺，已核对面积及批量",
                       "adjustments": [{"amount": -20, "reason": "材料价下降", "source": "采购价版本对照"}]}
    result = rfq.calculate(data)
    assert result["history_cost"] == Decimal("980.00")
    assert result["deviation"] == Decimal("0.25")/Decimal("980")
    data["history"]["actual_cost"] = 700
    assert any("超过 8%" in x for x in rfq.calculate(data)["reasons"])


def test_urgent_markup_and_missing_price_are_rejected(data):
    data["tiers"][1]["markup"] = .5
    with pytest.raises(ValueError, match="加急"): rfq.calculate(data)
    data["tiers"][1]["markup"] = .28
    data["cost_lines"][0]["price"] = None
    with pytest.raises(ValueError): rfq.calculate(data)


def test_excel_delivers_formulas_caches_and_literal_user_text(data, tmp_path):
    data["customer"] = '=HYPERLINK("https://invalid.test","customer")'
    result = rfq.calculate(data)
    files = rfq.export(data, result, tmp_path)
    assert len(files) == 4
    costs = openpyxl.load_workbook(files[0], data_only=True)
    formulas = openpyxl.load_workbook(files[0], data_only=False)
    sheet = costs["成本明细"]
    total_row = next(row[0].row for row in sheet.iter_rows() if row[1].value == "单台总成本")
    assert sheet.cell(total_row,6).value == 980.25
    assert formulas["成本明细"]["C2"].data_type == "f"
    assert formulas["报价状态"]["B3"].data_type == "s"
    quote = openpyxl.load_workbook(files[1], data_only=True)["三档报价"]
    quote_f = openpyxl.load_workbook(files[1], data_only=False)["三档报价"]
    for row,t in enumerate(result["tiers"],2):
        assert quote.cell(row,8).value == float(t["price"])
        assert quote.cell(row,10).value == float(t["profit"])
        assert quote_f.cell(row,8).value == f"=ROUND(F{row}*(1+G{row}),2)"
    risk = openpyxl.load_workbook(files[2], data_only=True)["风险明细"]
    assert risk["E7"].value == "待评估"
    with pytest.raises(ValueError, match="已有"): rfq.export(data,result,tmp_path)


def test_changed_material_reprices_all_tiers(data):
    old = rfq.calculate(data)
    data["material"]["price"] = 5.68
    new = rfq.calculate(data)
    assert new["cost"] > old["cost"]
    assert all(a["price"] > b["price"] for a,b in zip(new["tiers"],old["tiers"]))
