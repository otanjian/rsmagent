import importlib.util
import json
import stat
import sys
import zipfile
from pathlib import Path

import pytest

from test_rfq_agent_install import environment as identity_environment
from agent.skills.loader import SkillLoader
from scripts.build_workbuddy_presets import archive_files
from scripts.install_workbuddy_agents import PRESET, ensure_runtime, install_one, load_catalog, preflight, validate_bundle

spec = importlib.util.spec_from_file_location("scenario_io", PRESET/"shared/scenario_io.py")
io = importlib.util.module_from_spec(spec)
spec.loader.exec_module(io)
ENTRIES = load_catalog()


@pytest.fixture
def environment(identity_environment, monkeypatch):
    # Package installation is checked by the live smoke; unit tests never
    # download dependencies or create a separate runtime for each temp tenant.
    monkeypatch.setattr("scripts.install_workbuddy_agents.ensure_runtime", lambda *args: sys.executable)
    return identity_environment


def test_complete_catalog_and_original_assets_are_valid():
    assert len(ENTRIES) == 62
    assert {g: sum(e["group"] == g for e in ENTRIES) for g in ("manufacturing", "hr", "finance", "sales")} == {
        "manufacturing": 20, "hr": 15, "finance": 12, "sales": 15}
    assert sum(e["demo_file_count"] for e in ENTRIES) == 292
    for entry in ENTRIES:
        assert len(validate_bundle(entry)) == 64
        source = PRESET/"skills"/entry["skill"]
        scenario = json.loads((source/"scenario.json").read_text())
        assert len(scenario["demo_files"]) == entry["demo_file_count"]
        for name in scenario["demo_files"]:
            evidence = io.inspect_file(source/"assets/demo"/name)
            assert any(evidence.get(k) for k in ("sheets", "paragraphs", "tables", "pages", "lines")), name


def test_all_agents_install_with_private_skills_and_narrow_grants(environment):
    admin, svc, tenant, actor, other = environment
    default_agent = svc.tenant_default_agent_id(tenant)
    other_agents = svc.tenant_agent_ids(other)
    old_role = next(r for r in svc.list_roles(tenant) if r["code"] == "tenant_admin")
    old_other_roles = svc.list_roles(other)
    for entry in ENTRIES:
        result = install_one(admin, svc, entry, tenant_id=tenant, actor_user_id=actor)
        assert result["status"] == "installed"
        assert svc.get_agent_binding(result["id"])["tenant_id"] == tenant
        assert svc._member_can_use_agent(actor, tenant, result["id"])
        assert "custom:"+entry["skill"] in svc.resource_ids_for(actor, tenant, "skill", "use", permission="skill.use")
        workspace = Path(result["workspace"])
        assert workspace.is_relative_to(Path(svc.tenant_shared_root(tenant)))
        assert "UNRELATED PRIVATE PROFILE" not in (workspace/"USER.md").read_text()
        assert not (workspace/"BOOTSTRAP.md").exists()
        skills = SkillLoader().load_skills_from_dir(str(workspace/"skills"), "custom").skills
        assert [s.name for s in skills] == [entry["skill"]]
        assert (workspace/"knowledge").is_dir()
    role = next(r for r in svc.list_roles(tenant) if r["code"] == "tenant_admin")
    assert role["permissions"] == old_role["permissions"]
    assert role["model_defaults"] == old_role["model_defaults"]
    assert all(g in role["resource_grants"] for g in old_role["resource_grants"])
    assert len(role["resource_grants"])-len(old_role["resource_grants"]) == 62*4
    assert svc.list_roles(other) == old_other_roles
    assert svc.tenant_agent_ids(other) == other_agents
    assert svc.tenant_default_agent_id(tenant) == default_agent


def test_retry_preserves_customizations_and_repairs_only_missing_grant(environment):
    admin, svc, tenant, actor, _ = environment
    entry = ENTRIES[0]
    first = install_one(admin, svc, entry, tenant_id=tenant, actor_user_id=actor)
    path = Path(first["workspace"])/"AGENT.md"
    path.write_text("Tenant customization")
    role = next(r for r in svc.list_roles(tenant) if r["code"] == "tenant_admin")
    grants = [g for g in role["resource_grants"] if g["resource_id"] != "custom:"+entry["skill"]]
    svc.update_role(actor, tenant, role["id"], role["name"], role["permissions"], role["version"], resource_grants=grants)
    second = install_one(admin, svc, entry, tenant_id=tenant, actor_user_id=actor)
    assert second["status"] == "already_installed"
    assert path.read_text() == "Tenant customization"
    assert "custom:"+entry["skill"] in svc.resource_ids_for(actor, tenant, "skill", "use", permission="skill.use")


@pytest.mark.parametrize("operation", ["bind_agent", "update_role"])
def test_failure_compensates_new_agent_without_touching_existing(environment, monkeypatch, operation):
    admin, svc, tenant, actor, _ = environment
    good = install_one(admin, svc, ENTRIES[0], tenant_id=tenant, actor_user_id=actor)
    before = admin.snapshot()
    def fail(*args, **kwargs):
        raise RuntimeError("injected failure")
    monkeypatch.setattr(svc, operation, fail)
    with pytest.raises(RuntimeError, match="injected"):
        install_one(admin, svc, ENTRIES[1], tenant_id=tenant, actor_user_id=actor)
    assert [a["id"] for a in admin.snapshot()["agents"]] == [a["id"] for a in before["agents"]]
    assert svc.get_agent_binding(good["id"])["tenant_id"] == tenant
    assert not (Path(good["workspace"]).parent/(ENTRIES[1]["skill"]+"-test15")).exists()


def test_cross_tenant_and_occupied_workspace_rejected(environment):
    admin, svc, tenant, actor, other = environment
    with pytest.raises(ValueError, match="管理员"):
        preflight(admin, svc, ENTRIES[0], tenant_id=other, actor_user_id=actor)
    workspace = Path(svc.tenant_shared_root(tenant))/"agents"/(ENTRIES[0]["skill"]+"-test15")
    workspace.mkdir(parents=True)
    sentinel = workspace/"user.txt"
    sentinel.write_text("keep")
    with pytest.raises(ValueError, match="工作区已存在"):
        install_one(admin, svc, ENTRIES[0], tenant_id=tenant, actor_user_id=actor)
    assert sentinel.read_text() == "keep"


@pytest.mark.parametrize("kind", ["traversal", "absolute", "symlink", "executable"])
def test_archive_import_rejects_unsafe_members(tmp_path, kind):
    archive = tmp_path/"pkg.zip"
    with zipfile.ZipFile(archive, "w") as z:
        z.writestr("SKILL.md", "business rules")
        z.writestr("输入参考文件/demo.txt", "data")
        name = {"traversal": "../x", "absolute": "/tmp/x", "symlink": "references/link.md", "executable": "scripts/run.sh"}[kind]
        info = zipfile.ZipInfo(name)
        if kind == "symlink":
            info.external_attr = (stat.S_IFLNK | 0o777) << 16
        z.writestr(info, "payload")
    with pytest.raises(ValueError):
        archive_files(archive)


def test_prepare_all_demos_is_explicit_non_destructive_and_not_a_fake_result(tmp_path):
    for entry in ENTRIES:
        source = PRESET/"skills"/entry["skill"]
        out = tmp_path/entry["slug"]
        prepared = io.prepare(source, out)
        assert prepared["state"] == "inputs_prepared"
        assert len(prepared["files"]) == entry["demo_file_count"]
        assert not list((out/"outputs").iterdir())
        for record in prepared["files"]:
            path = Path(record["path"])
            assert path.read_bytes() == (source/"assets/demo"/path.name).read_bytes()
        with pytest.raises(FileExistsError):
            io.prepare(source, out)


def test_render_real_files_preserves_formulas_and_does_not_execute_input_text(tmp_path):
    import openpyxl
    spec = {"title": "核对报告", "basis": "合成测试；来源 A2:C2", "files": [
        {"filename": "核对.xlsx", "sheets": [{"name": "明细", "columns": ["数量", "单价", "金额", "原文"],
          "rows": [[2, 3, {"formula": "=A4*B4", "value": 6}, "=malicious()"]]}]},
        {"filename": "说明.pdf", "paragraphs": ["核对结果为六元。待审批，不自动发出。"]},
        {"filename": "说明.docx", "paragraphs": ["核对结果为六元。"]},
        {"filename": "待办.md", "text": "请核实数量来源。"}]}
    result = io.render(spec, tmp_path)
    assert len(result) == 4 and all(r["readable"] for r in result)
    book = openpyxl.load_workbook(tmp_path/"核对.xlsx", data_only=False)
    assert book["明细"]["C4"].value == "=A4*B4"
    assert book["明细"]["D4"].data_type == "s"
    book.close()
    book = openpyxl.load_workbook(tmp_path/"核对.xlsx", data_only=True)
    assert book["明细"]["C4"].value == 6
    book.close()
    assert "核对结果" in io.inspect_file(tmp_path/"说明.pdf")["pages"][0]["text"]
    with pytest.raises(ValueError, match="已存在"):
        io.render(spec, tmp_path)


def test_render_failure_removes_only_new_partial_files(tmp_path):
    sentinel = tmp_path/"existing.txt"
    sentinel.write_text("keep")
    spec = {"title": "t", "basis": "b", "files": [{"filename": "new.md", "text": "data"},
            {"filename": "bad.xlsx", "sheets": [{"name": "x", "columns": ["a"], "rows": [[{"formula": "=1+2"}]]}]}]}
    with pytest.raises(ValueError, match="公式"):
        io.render(spec, tmp_path)
    assert list(tmp_path.iterdir()) == [sentinel]


def test_check_rejects_formula_errors_and_render_path_escape(tmp_path):
    import openpyxl
    book = openpyxl.Workbook()
    book.active["A1"] = "#REF!"
    path = tmp_path/"error.xlsx"
    book.save(path)
    with pytest.raises(ValueError, match="#REF!"):
        io.check(path)
    with pytest.raises(ValueError, match="单个文件名"):
        io.render({"title": "t", "basis": "b", "files": [{"filename": "../escape.md", "text": "x"}]}, tmp_path)


def test_runtime_refuses_existing_directory_and_external_interpreter(tmp_path):
    import hashlib
    runtime = tmp_path/"runtimes/workbuddy"
    runtime.mkdir(parents=True)
    sentinel = runtime/"user.txt"
    sentinel.write_text("keep")
    with pytest.raises(ValueError, match="占用"):
        ensure_runtime(tmp_path)
    assert sentinel.read_text() == "keep"
    marker = {"version": 1, "requirements_sha256": hashlib.sha256((PRESET/"requirements.txt").read_bytes()).hexdigest()}
    (runtime/"workbuddy-runtime.json").write_text(json.dumps(marker))
    executable = runtime/"env/bin/python"
    executable.parent.mkdir(parents=True)
    executable.symlink_to(sys.executable)
    with pytest.raises(ValueError, match="解释器必须位于租户"):
        ensure_runtime(tmp_path)
    assert sentinel.read_text() == "keep"
